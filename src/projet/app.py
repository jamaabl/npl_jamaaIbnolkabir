"""Détecteur de Plagiat (Corpus) - interface customtkinter.

Dépendances :  pip install customtkinter
Optionnel   :  pip install pypdf python-docx   (lecture des PDF et DOCX)
Ce fichier utilise tes modules existants :
    preprocess.py -> clean_text(texte), remove_stopwords(texte)
    model.py      -> detect_plagiarism(segments_suspects, segments_source, threshold)
                     doit renvoyer une liste de dict avec au moins 'score' et 'suspicious'
"""
import datetime
import glob
import json
import os
import queue
import threading
import time
import customtkinter as ctk
from tkinter import filedialog, messagebox

from preprocess import clean_text, remove_stopwords
from model import detect_plagiarism

ctk.set_appearance_mode("Light")
ctk.set_default_color_theme("blue")

# ---------- Couleurs ----------
NAVY, NAVY_DARK = "#0b2a5c", "#081f45"
BLUE, BLUE_HOVER = "#1463e6", "#0f52c4"
BG, CARD, BORDER = "#f3f6fb", "#ffffff", "#dce4f0"
TEXT, MUTED = "#0f2956", "#6b7a90"
RED, GREEN = "#e74c3c", "#27ae60"

HISTO_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "historique.json")
EXTENSIONS = (".txt", ".pdf", ".docx", ".xlsx", ".xls")


# ---------- Lecture de fichiers ----------
def lire_fichier(chemin):
    """Extrait le texte d'un document : PDF, Word (DOCX), Excel (XLSX/XLS) ou TXT."""
    ext = os.path.splitext(chemin)[1].lower()
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            raise RuntimeError("Pour lire les PDF : pip install pypdf  (ou uv add pypdf)")
        return "\n".join((p.extract_text() or "") for p in PdfReader(chemin).pages)
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            raise RuntimeError("Pour lire les DOCX : pip install python-docx  (ou uv add python-docx)")
        d = docx.Document(chemin)
        morceaux = [p.text for p in d.paragraphs]
        for t in d.tables:                       # texte des tableaux Word
            for row in t.rows:
                morceaux.extend(c.text for c in row.cells)
        return "\n".join(morceaux)
    if ext in (".xlsx", ".xls"):
        try:
            import pandas as pd
        except ImportError:
            raise RuntimeError("Pour lire les Excel : pip install pandas openpyxl  (ou uv add pandas openpyxl)")
        morceaux = []
        for df in pd.read_excel(chemin, sheet_name=None, header=None).values():   # toutes les feuilles
            for v in df.astype(str).values.ravel():
                if v and v.lower() != "nan":
                    morceaux.append(v)
        return "\n".join(morceaux)
    with open(chemin, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def segmenter(texte, taille=300):
    return [texte[i:i + taille] for i in range(0, len(texte), taille)]


def fmt_temps(sec):
    sec = int(sec)
    return f"{sec // 60}:{sec % 60:02d}"


class AppModerne(ctk.CTk):
    def __init__(self):
        super().__init__()
        self.title("Détecteur de Plagiat (Corpus)")
        self.geometry("1280x860")
        self.minsize(1100, 760)
        self.configure(fg_color=BG)

        # état
        self.dossier_corpus = ""
        self.nom_document = ""      # nom du document suspect chargé
        self.contenu_suspect = ""   # texte extrait du document
        self.doc_analyse = ""
        self.en_cours = False
        self.stop_event = threading.Event()
        self.file_q = queue.Queue()
        self.t_debut = 0.0
        self.historique = self._charger_historique()
        # paramètres
        self.var_seuil = ctk.DoubleVar(value=0.5)
        self.var_max = ctk.StringVar(value="100")
        self.var_taille = ctk.StringVar(value="300")

        self.grid_columnconfigure(1, weight=1)
        self.grid_rowconfigure(1, weight=1)
        self._build_sidebar()
        self._build_header()
        self._build_pages()
        self._build_statusbar()
        self.afficher_page("accueil")
        try:
            self.after(100, lambda: self.state("zoomed"))   # plein écran sous Windows
        except Exception:
            pass

    # =====================================================================
    #  CONSTRUCTION DE L'INTERFACE
    # =====================================================================
    def _build_sidebar(self):
        sb = ctk.CTkFrame(self, width=210, corner_radius=0, fg_color=NAVY)
        sb.grid(row=0, column=0, rowspan=3, sticky="nsew")
        sb.pack_propagate(False)
        ctk.CTkLabel(sb, text="🛡️ Détecteur\nde Plagiat", font=("Arial", 20, "bold"),
                     text_color="white").pack(pady=(30, 25), padx=20)
        self.nav = {}
        for cle, texte in [("accueil", "🏠  Accueil"), ("analyse", "📄  Analyse"),
                           ("historique", "🕘  Historique"), ("parametres", "⚙️  Paramètres")]:
            b = ctk.CTkButton(sb, text=texte, anchor="w", height=42, corner_radius=8,
                              fg_color="transparent", hover_color="#17407f", font=("Arial", 14),
                              command=lambda c=cle: self.afficher_page(c))
            b.pack(fill="x", padx=12, pady=4)
            self.nav[cle] = b
        ctk.CTkLabel(sb, text="ℹ️ Détecteur de Plagiat v1.0\nVotre confiance,\nnotre priorité",
                     text_color="#aebcd6", font=("Arial", 11), justify="center").pack(side="bottom", pady=20)

    def _build_header(self):
        h = ctk.CTkFrame(self, height=100, corner_radius=0, fg_color=NAVY_DARK)
        h.grid(row=0, column=1, sticky="ew")
        h.grid_propagate(False)
        h.grid_columnconfigure(0, weight=1)
        h.grid_rowconfigure(0, weight=1)
        g = ctk.CTkFrame(h, fg_color="transparent")
        g.grid(row=0, column=0, sticky="w", padx=30)
        ctk.CTkLabel(g, text="Détecteur de Plagiat (Corpus)", font=("Arial", 26, "bold"),
                     text_color="white").pack(anchor="w")
        ctk.CTkLabel(g, text="Analysez et vérifiez l'originalité de vos documents",
                     font=("Arial", 14), text_color="#b7c7e6").pack(anchor="w")
        ctk.CTkLabel(h, text="🧠  NLP  •  Analyse sémantique  •  Précision", font=("Arial", 12),
                     text_color="#b7c7e6").grid(row=0, column=1, padx=30)

    def _build_pages(self):
        self.zone = ctk.CTkFrame(self, fg_color=BG, corner_radius=0)
        self.zone.grid(row=1, column=1, sticky="nsew")
        self.zone.grid_columnconfigure(0, weight=1)
        self.zone.grid_rowconfigure(0, weight=1)
        self.pages = {}
        for cle in ("accueil", "analyse", "historique", "parametres"):
            f = ctk.CTkFrame(self.zone, fg_color=BG, corner_radius=0)
            f.grid(row=0, column=0, sticky="nsew")
            self.pages[cle] = f
        self._page_accueil(self.pages["accueil"])
        self._page_analyse(self.pages["analyse"])
        self._page_historique(self.pages["historique"])
        self._page_parametres(self.pages["parametres"])

    def _build_statusbar(self):
        s = ctk.CTkFrame(self, height=40, corner_radius=0, fg_color=NAVY_DARK)
        s.grid(row=2, column=1, sticky="ew")
        s.grid_propagate(False)
        s.grid_columnconfigure(1, weight=1)
        s.grid_rowconfigure(0, weight=1)
        self.dot = ctk.CTkLabel(s, text="●", text_color=GREEN, font=("Arial", 14))
        self.dot.grid(row=0, column=0, padx=(20, 4))
        self.lbl_statut = ctk.CTkLabel(s, text="Prêt", text_color="white", font=("Arial", 12))
        self.lbl_statut.grid(row=0, column=1, sticky="w")
        ctk.CTkLabel(s, text="Détecteur de Plagiat (Corpus)  |  NLP & IA", text_color="#b7c7e6",
                     font=("Arial", 12)).grid(row=0, column=2, padx=20)

    # ---------- helpers de style ----------
    def _carte(self, parent):
        return ctk.CTkFrame(parent, fg_color=CARD, corner_radius=12, border_width=1, border_color=BORDER)

    def _titre_carte(self, parent, icone, titre, sous_titre=None):
        f = ctk.CTkFrame(parent, fg_color="transparent")
        f.pack(fill="x", padx=18, pady=(16, 6))
        ctk.CTkLabel(f, text=icone, width=44, height=44, corner_radius=22, fg_color=BLUE,
                     text_color="white", font=("Arial", 18)).pack(side="left")
        col = ctk.CTkFrame(f, fg_color="transparent")
        col.pack(side="left", padx=12)
        ctk.CTkLabel(col, text=titre, font=("Arial", 16, "bold"), text_color=TEXT).pack(anchor="w")
        if sous_titre:
            ctk.CTkLabel(col, text=sous_titre, font=("Arial", 12), text_color=MUTED,
                         wraplength=520, justify="left").pack(anchor="w")
        return f

    def _btn_contour(self, parent, texte, commande, couleur=BLUE):
        return ctk.CTkButton(parent, text=texte, command=commande, fg_color="white", text_color=couleur,
                             border_width=1, border_color=couleur, hover_color="#eef3fb")

    # ---------- PAGE ACCUEIL ----------
    def _page_accueil(self, page):
        page.grid_columnconfigure(0, weight=3, uniform="c")
        page.grid_columnconfigure(1, weight=2, uniform="c")
        page.grid_rowconfigure(0, weight=1)
        gauche = ctk.CTkFrame(page, fg_color="transparent")
        gauche.grid(row=0, column=0, sticky="nsew", padx=(24, 8), pady=16)
        droite = ctk.CTkFrame(page, fg_color="transparent")
        droite.grid(row=0, column=1, sticky="nsew", padx=(8, 24), pady=16)
        gauche.grid_columnconfigure(0, weight=1)
        gauche.grid_rowconfigure(0, weight=1)
        droite.grid_columnconfigure(0, weight=1)
        droite.grid_rowconfigure(0, weight=3)
        droite.grid_rowconfigure(2, weight=2)

        # --- 1. Document suspect ---
        c1 = self._carte(gauche)
        c1.grid(row=0, column=0, sticky="nsew", pady=(0, 12))
        self._titre_carte(c1, "📄", "1. Document suspect",
                          "Chargez le document à analyser (PDF, Word, Excel ou TXT).")
        self.lbl_doc = ctk.CTkLabel(c1, text="Aucun document chargé", anchor="w", height=40,
                                    fg_color="#eef3fa", corner_radius=8, text_color=MUTED,
                                    font=("Arial", 13, "bold"))
        self.lbl_doc.pack(fill="x", padx=18, pady=(6, 6))
        # boutons juste sous le nom du document : toujours visibles, même fenêtre basse
        barre = ctk.CTkFrame(c1, fg_color="transparent")
        barre.pack(fill="x", padx=18, pady=(0, 6))
        self.btn_charger = ctk.CTkButton(barre, text="⬆  Charger un document", height=38,
                                         command=self.charger_fichier)
        self.btn_charger.pack(side="left")
        self.btn_effacer = self._btn_contour(barre, "🗑  Retirer le document", self.effacer_texte, couleur=TEXT)
        self.btn_effacer.configure(height=38)
        self.btn_effacer.pack(side="right")
        self.lbl_compteur = ctk.CTkLabel(c1, text="0 caractère", text_color=MUTED, font=("Arial", 11))
        self.lbl_compteur.pack(side="bottom", anchor="e", padx=18, pady=(0, 10))
        ctk.CTkLabel(c1, text="Aperçu du texte extrait (lecture seule)", anchor="w", text_color=MUTED,
                     font=("Arial", 11)).pack(fill="x", padx=18)
        self.txt_suspect = ctk.CTkTextbox(c1, font=("Arial", 12), height=80, fg_color="#fafbfd",
                                          border_width=1, border_color=BORDER, text_color=TEXT,
                                          state="disabled")
        self.txt_suspect.pack(fill="both", expand=True, padx=18, pady=(4, 4))

        # --- 2. Corpus + lancement ---
        c2 = self._carte(gauche)
        c2.grid(row=1, column=0, sticky="ew")
        self._titre_carte(c2, "📁", "2. Sélectionner le dossier du Corpus",
                          "Choisissez le dossier contenant les documents de référence (corpus) à comparer.")
        ligne = ctk.CTkFrame(c2, fg_color="transparent")
        ligne.pack(fill="x", padx=18, pady=(6, 6))
        self.lbl_dossier = ctk.CTkLabel(ligne, text="  Aucun dossier sélectionné", anchor="w", height=38,
                                        fg_color="#eef3fa", corner_radius=6, text_color=MUTED)
        self.lbl_dossier.pack(side="left", fill="x", expand=True, padx=(0, 10))
        self.btn_parcourir = ctk.CTkButton(ligne, text="📂  Parcourir...", width=130, height=38,
                                           command=self.choisir_dossier)
        self.btn_parcourir.pack(side="right")
        ctk.CTkLabel(c2, text="ℹ️  Le corpus doit contenir des fichiers texte (TXT, PDF, DOCX, XLSX) dans un dossier.",
                     fg_color="#eef3fa", corner_radius=6, text_color=TEXT, font=("Arial", 11),
                     anchor="w", height=34).pack(fill="x", padx=18, pady=(2, 8))
        self.btn_lancer = ctk.CTkButton(c2, text="▶  Lancer la Détection sur le Corpus", height=52,
                                        font=("Arial", 16, "bold"), command=self.basculer_analyse)
        self.btn_lancer.pack(fill="x", padx=60, pady=(6, 10))
        pf = ctk.CTkFrame(c2, fg_color="#f7f9fc", border_width=1, border_color=BORDER, corner_radius=8)
        pf.pack(fill="x", padx=18, pady=(0, 16))
        self.lbl_etat = ctk.CTkLabel(pf, text="En attente d'une analyse", font=("Arial", 13, "bold"),
                                     text_color=TEXT)
        self.lbl_etat.pack(anchor="w", padx=14, pady=(10, 0))
        self.lbl_etat_sub = ctk.CTkLabel(pf, text="Chargez un document, choisissez un corpus, puis lancez.",
                                         font=("Arial", 11), text_color=MUTED)
        self.lbl_etat_sub.pack(anchor="w", padx=14)
        bp = ctk.CTkFrame(pf, fg_color="transparent")
        bp.pack(fill="x", padx=14, pady=(4, 12))
        self.progress = ctk.CTkProgressBar(bp)
        self.progress.pack(side="left", fill="x", expand=True)
        self.progress.set(0)
        self.lbl_pct = ctk.CTkLabel(bp, text="0%", width=45, text_color=MUTED)
        self.lbl_pct.pack(side="right", padx=(8, 0))

        # --- Résultats ---
        c3 = self._carte(droite)
        c3.grid(row=0, column=0, sticky="nsew", pady=(0, 12))
        top = ctk.CTkFrame(c3, fg_color="transparent")
        top.pack(fill="x", padx=18, pady=(16, 6))
        ctk.CTkLabel(top, text="📊", width=40, height=40, corner_radius=20, fg_color="#e3edff",
                     font=("Arial", 16)).pack(side="left")
        ctk.CTkLabel(top, text="Résultats de l'analyse", font=("Arial", 16, "bold"),
                     text_color=TEXT).pack(side="left", padx=10)
        self.badge = ctk.CTkLabel(top, text="Aucun résultat", fg_color="#e3edff", text_color=BLUE,
                                  corner_radius=10, font=("Arial", 11), width=110, height=24)
        self.badge.pack(side="right")
        body = ctk.CTkFrame(c3, fg_color="transparent")
        body.pack(fill="both", expand=True, padx=18, pady=(0, 16))
        self.placeholder = ctk.CTkFrame(body, fg_color="#fbfcfe", border_width=1, border_color=BORDER)
        ctk.CTkLabel(self.placeholder, text="📄🔍", font=("Arial", 28)).pack(pady=(14, 4))
        ctk.CTkLabel(self.placeholder, text="Aucun résultat pour le moment", font=("Arial", 14, "bold"),
                     text_color=TEXT).pack()
        ctk.CTkLabel(self.placeholder, text="Lancez une analyse pour découvrir\nles similitudes éventuelles.",
                     font=("Arial", 12), text_color=MUTED).pack(pady=(2, 12))
        self.txt_resultats = ctk.CTkTextbox(body, font=("Consolas", 12), fg_color="#f9f9f9",
                                            text_color="#333", border_width=1, border_color=BORDER)
        self._montrer_placeholder(True)

        # --- Statistiques ---
        c4 = self._carte(droite)
        c4.grid(row=1, column=0, sticky="ew", pady=(0, 12))
        ctk.CTkLabel(c4, text="📈  Statistiques", font=("Arial", 15, "bold"),
                     text_color=TEXT).pack(anchor="w", padx=18, pady=(14, 6))
        grille = ctk.CTkFrame(c4, fg_color="transparent")
        grille.pack(fill="x", padx=14, pady=(0, 14))
        grille.grid_columnconfigure((0, 1), weight=1, uniform="s")
        self.val_fichiers = self._tuile(grille, 0, 0, "📄 Fichiers analysés", "0")
        self.val_compar = self._tuile(grille, 0, 1, "🔍 Comparaison avec la base", "0%")
        self.val_plagiats = self._tuile(grille, 1, 0, "🛡️ Plagiats détectés", "0")
        self.val_temps = self._tuile(grille, 1, 1, "⏱️ Temps écoulé", "0:00")

        # --- Fichiers similaires ---
        c5 = self._carte(droite)
        c5.grid(row=2, column=0, sticky="nsew")
        ctk.CTkLabel(c5, text="📑  Fichiers similaires", font=("Arial", 15, "bold"),
                     text_color=TEXT).pack(anchor="w", padx=18, pady=(14, 4))
        self.liste_sim = ctk.CTkScrollableFrame(c5, fg_color="transparent", height=110)
        self.liste_sim.pack(fill="both", expand=True, padx=10, pady=(0, 10))
        self._vider_similaires()

    def _tuile(self, parent, r, c, titre, valeur):
        t = ctk.CTkFrame(parent, fg_color="#f7f9fc", border_width=1, border_color=BORDER, corner_radius=8)
        t.grid(row=r, column=c, sticky="nsew", padx=4, pady=4)
        ctk.CTkLabel(t, text=titre, font=("Arial", 11), text_color=MUTED).pack(anchor="w", padx=12, pady=(8, 0))
        lbl = ctk.CTkLabel(t, text=valeur, font=("Arial", 22, "bold"), text_color=TEXT)
        lbl.pack(anchor="w", padx=12, pady=(0, 8))
        return lbl

    # ---------- PAGE ANALYSE ----------
    def _page_analyse(self, page):
        page.grid_columnconfigure(0, weight=1)
        page.grid_rowconfigure(0, weight=1)
        c = self._carte(page)
        c.grid(row=0, column=0, sticky="nsew", padx=24, pady=16)
        self._titre_carte(c, "📄", "Rapport détaillé de la dernière analyse",
                          "Toutes les correspondances (jusqu'à 50) avec le passage suspect.")
        self.txt_rapport = ctk.CTkTextbox(c, font=("Consolas", 12), fg_color="#f9f9f9", text_color="#333",
                                          border_width=1, border_color=BORDER)
        self.txt_rapport.pack(fill="both", expand=True, padx=18, pady=(6, 8))
        self.txt_rapport.insert("1.0", "Aucune analyse effectuée pour le moment.")
        barre = ctk.CTkFrame(c, fg_color="transparent")
        barre.pack(fill="x", padx=18, pady=(0, 16))
        ctk.CTkButton(barre, text="💾  Exporter le rapport (.txt)", command=self.exporter_rapport).pack(side="left")
        self._btn_contour(barre, "🏠  Retour à l'accueil", lambda: self.afficher_page("accueil")).pack(side="right")

    # ---------- PAGE HISTORIQUE ----------
    def _page_historique(self, page):
        page.grid_columnconfigure(0, weight=1)
        page.grid_rowconfigure(0, weight=1)
        c = self._carte(page)
        c.grid(row=0, column=0, sticky="nsew", padx=24, pady=16)
        self._titre_carte(c, "🕘", "Historique des analyses", "Cliquez sur « Voir » pour rouvrir un rapport.")
        self.liste_histo = ctk.CTkScrollableFrame(c, fg_color="transparent")
        self.liste_histo.pack(fill="both", expand=True, padx=12, pady=(6, 8))
        barre = ctk.CTkFrame(c, fg_color="transparent")
        barre.pack(fill="x", padx=18, pady=(0, 16))
        ctk.CTkButton(barre, text="🗑  Vider l'historique", fg_color=RED, hover_color="#c0392b",
                      command=self.vider_historique).pack(side="right")

    # ---------- PAGE PARAMÈTRES ----------
    def _page_parametres(self, page):
        page.grid_columnconfigure(0, weight=1)
        c = self._carte(page)
        c.grid(row=0, column=0, sticky="new", padx=24, pady=16)
        self._titre_carte(c, "⚙️", "Paramètres de détection", "Ces valeurs sont appliquées à la prochaine analyse.")
        f = ctk.CTkFrame(c, fg_color="transparent")
        f.pack(fill="x", padx=24, pady=(8, 8))
        f.grid_columnconfigure(1, weight=1)

        ctk.CTkLabel(f, text="Seuil de similarité", text_color=TEXT, font=("Arial", 13, "bold")).grid(
            row=0, column=0, sticky="w", pady=12)
        self.lbl_seuil = ctk.CTkLabel(f, text=f"{self.var_seuil.get():.2f}", text_color=BLUE,
                                      font=("Arial", 14, "bold"), width=50)
        self.lbl_seuil.grid(row=0, column=2, padx=10)
        ctk.CTkSlider(f, from_=0.1, to=1.0, number_of_steps=18, variable=self.var_seuil,
                      command=lambda v: self.lbl_seuil.configure(text=f"{float(v):.2f}")).grid(
            row=0, column=1, sticky="ew", padx=20)

        ctk.CTkLabel(f, text="Nombre max de fichiers (0 = tous)", text_color=TEXT,
                     font=("Arial", 13, "bold")).grid(row=1, column=0, sticky="w", pady=12)
        ctk.CTkEntry(f, textvariable=self.var_max, width=100).grid(row=1, column=1, sticky="w", padx=20)

        ctk.CTkLabel(f, text="Taille d'un segment (caractères)", text_color=TEXT,
                     font=("Arial", 13, "bold")).grid(row=2, column=0, sticky="w", pady=12)
        ctk.CTkEntry(f, textvariable=self.var_taille, width=100).grid(row=2, column=1, sticky="w", padx=20)

        ctk.CTkLabel(c, text="Plus le seuil est élevé, plus la détection est stricte (moins de faux positifs).",
                     text_color=MUTED, font=("Arial", 11)).pack(anchor="w", padx=24)
        ctk.CTkButton(c, text="↺  Valeurs par défaut", command=self.reinitialiser_parametres).pack(
            anchor="w", padx=24, pady=(10, 20))

    # =====================================================================
    #  NAVIGATION / UTILITAIRES UI
    # =====================================================================
    def afficher_page(self, cle):
        self.pages[cle].tkraise()
        for k, b in self.nav.items():
            b.configure(fg_color=BLUE if k == cle else "transparent")
        if cle == "historique":
            self._rafraichir_historique()

    def _set_statut(self, texte, couleur=GREEN):
        self.lbl_statut.configure(text=texte)
        self.dot.configure(text_color=couleur)

    def _afficher_apercu(self):
        contenu = self.contenu_suspect
        self.txt_suspect.configure(state="normal")
        self.txt_suspect.delete("1.0", "end")
        self.txt_suspect.insert("1.0", contenu[:5000] + ("\n[...]" if len(contenu) > 5000 else ""))
        self.txt_suspect.configure(state="disabled")
        if self.nom_document:
            self.lbl_doc.configure(text=f"  📎  {self.nom_document}", text_color=BLUE)
            self.lbl_compteur.configure(text=f"{len(contenu)} caractères · {len(contenu.split())} mots extraits")
        else:
            self.lbl_doc.configure(text="Aucun document chargé", text_color=MUTED)
            self.lbl_compteur.configure(text="0 caractère")

    def _montrer_placeholder(self, oui):
        if oui:
            self.txt_resultats.pack_forget()
            self.placeholder.pack(fill="both", expand=True)
        else:
            self.placeholder.pack_forget()
            self.txt_resultats.pack(fill="both", expand=True)

    def _vider_similaires(self):
        for w in self.liste_sim.winfo_children():
            w.destroy()
        ctk.CTkLabel(self.liste_sim, text="Aucun fichier similaire trouvé\nLes résultats s'afficheront ici "
                     "après l'analyse.", text_color=MUTED, font=("Arial", 11)).pack(pady=14)

    def _etat_boutons(self, actifs):
        etat = "normal" if actifs else "disabled"
        for b in (self.btn_parcourir, self.btn_charger, self.btn_effacer):
            b.configure(state=etat)

    # =====================================================================
    #  ACTIONS DES BOUTONS
    # =====================================================================
    def charger_fichier(self):
        chemin = filedialog.askopenfilename(
            title="Charger le document suspect",
            filetypes=[("Documents supportés", "*.pdf *.docx *.xlsx *.xls *.txt"), ("PDF", "*.pdf"),
                       ("Word", "*.docx"), ("Excel", "*.xlsx *.xls"), ("Texte", "*.txt"),
                       ("Tous les fichiers", "*.*")])
        if not chemin:
            return
        try:
            contenu = lire_fichier(chemin)
        except Exception as e:
            messagebox.showerror("Erreur de lecture", str(e))
            return
        if not contenu.strip():
            messagebox.showwarning("Document vide", "Aucun texte n'a pu être extrait de ce document "
                                   "(PDF scanné / image ?).")
            return
        self.contenu_suspect = contenu
        self.nom_document = os.path.basename(chemin)
        self._afficher_apercu()

    def effacer_texte(self):
        self.contenu_suspect = ""
        self.nom_document = ""
        self._afficher_apercu()

    def choisir_dossier(self):
        d = filedialog.askdirectory(title="Sélectionnez le dossier du corpus")
        if d:
            self.dossier_corpus = d
            self.lbl_dossier.configure(text=f"  {d}", text_color=TEXT)

    def reinitialiser_parametres(self):
        self.var_seuil.set(0.5)
        self.lbl_seuil.configure(text="0.50")
        self.var_max.set("100")
        self.var_taille.set("300")

    def exporter_rapport(self):
        contenu = self.txt_rapport.get("1.0", "end").strip()
        chemin = filedialog.asksaveasfilename(defaultextension=".txt", initialfile="rapport_plagiat.txt",
                                              filetypes=[("Texte", "*.txt")])
        if chemin:
            with open(chemin, "w", encoding="utf-8") as f:
                f.write(contenu)
            messagebox.showinfo("Export", "Rapport enregistré :\n" + chemin)

    def ouvrir_fichier(self, chemin):
        try:
            os.startfile(chemin)
        except AttributeError:
            messagebox.showinfo("Fichier", chemin)
        except OSError as e:
            messagebox.showerror("Erreur", str(e))

    # ---------- historique ----------
    def _charger_historique(self):
        try:
            with open(HISTO_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except (OSError, ValueError):
            return []

    def _sauver_historique(self):
        try:
            with open(HISTO_FILE, "w", encoding="utf-8") as f:
                json.dump(self.historique[-50:], f, ensure_ascii=False, indent=1)
        except OSError:
            pass

    def _rafraichir_historique(self):
        for w in self.liste_histo.winfo_children():
            w.destroy()
        if not self.historique:
            ctk.CTkLabel(self.liste_histo, text="Aucune analyse enregistrée.", text_color=MUTED).pack(pady=30)
            return
        for e in reversed(self.historique):
            ligne = ctk.CTkFrame(self.liste_histo, fg_color="#f7f9fc", border_width=1, border_color=BORDER)
            ligne.pack(fill="x", pady=4, padx=4)
            txt = (f"{e['date']}  —  {e.get('document', 'texte saisi')}  —  {e['fichiers']} fichier(s) · {e['plagiats']} plagiat(s) · "
                   f"{e['temps']}\n« {e['extrait']}... »")
            ctk.CTkLabel(ligne, text=txt, justify="left", anchor="w", text_color=TEXT,
                         font=("Arial", 12)).pack(side="left", padx=12, pady=8)
            ctk.CTkButton(ligne, text="Voir", width=70,
                          command=lambda r=e["rapport"]: self._voir_rapport(r)).pack(side="right", padx=12)

    def _voir_rapport(self, rapport):
        self.txt_rapport.delete("1.0", "end")
        self.txt_rapport.insert("1.0", rapport)
        self.afficher_page("analyse")

    def vider_historique(self):
        if self.historique and messagebox.askyesno("Historique", "Supprimer tout l'historique ?"):
            self.historique = []
            self._sauver_historique()
            self._rafraichir_historique()

    # =====================================================================
    #  ANALYSE (thread séparé pour ne pas figer la fenêtre)
    # =====================================================================
    def basculer_analyse(self):
        if self.en_cours:
            self.stop_event.set()
            self.btn_lancer.configure(state="disabled", text="Arrêt en cours...")
        else:
            self.lancer_detection()

    def lancer_detection(self):
        texte = self.contenu_suspect.strip()
        if not texte:
            messagebox.showwarning("Attention", "Veuillez charger le document suspect (PDF, Word, Excel ou TXT).")
            return
        if not self.dossier_corpus:
            messagebox.showwarning("Attention", "Veuillez choisir le dossier du corpus.")
            return
        try:
            max_f, taille = int(self.var_max.get()), int(self.var_taille.get())
            if max_f < 0 or taille < 50:
                raise ValueError
        except ValueError:
            messagebox.showwarning("Paramètres", "Nombre de fichiers : entier >= 0.\n"
                                   "Taille de segment : entier >= 50 (voir page Paramètres).")
            return
        seuil = round(self.var_seuil.get(), 2)
        self.doc_analyse = self.nom_document

        # remise à zéro de l'affichage
        self.txt_resultats.delete("1.0", "end")
        self._montrer_placeholder(True)
        self.badge.configure(text="Analyse...", fg_color="#e3edff", text_color=BLUE)
        self._vider_similaires()
        self.progress.set(0)
        self.lbl_pct.configure(text="0%")
        for lbl, v in ((self.val_fichiers, "0"), (self.val_compar, "0%"),
                       (self.val_plagiats, "0"), (self.val_temps, "0:00")):
            lbl.configure(text=v)
        self.lbl_etat.configure(text="Analyse en cours...")
        self.lbl_etat_sub.configure(text="Veuillez patienter, le traitement du corpus peut prendre quelques instants.")
        self._set_statut("Analyse en cours...", "#f39c12")

        self.en_cours = True
        self.stop_event.clear()
        self.t_debut = time.time()
        self.btn_lancer.configure(text="⏹  Arrêter l'analyse", fg_color=RED, hover_color="#c0392b")
        self._etat_boutons(False)
        threading.Thread(target=self._worker, args=(texte, self.dossier_corpus, seuil, max_f, taille),
                         daemon=True).start()
        self.after(100, self._poll)

    def _worker(self, texte, dossier, seuil, max_f, taille):
        """Tourne dans un thread : ne touche PAS aux widgets, communique par file."""
        try:
            sus_seg = segmenter(remove_stopwords(clean_text(texte)), taille)
            if not sus_seg:
                self.file_q.put(("erreur", "Le texte suspect est vide après prétraitement."))
                return
            fichiers = sorted({f for ext in EXTENSIONS
                               for f in glob.glob(os.path.join(dossier, "**", "*" + ext), recursive=True)})
            if max_f > 0:
                fichiers = fichiers[:max_f]
            if not fichiers:
                self.file_q.put(("erreur", "Aucun fichier .txt/.pdf/.docx/.xlsx trouvé dans ce dossier."))
                return

            resultats, flagged, erreurs, interrompu = [], set(), 0, False
            total = len(fichiers)
            for i, f in enumerate(fichiers):
                if self.stop_event.is_set():
                    interrompu = True
                    break
                try:
                    contenu = lire_fichier(f)
                    if len(contenu) > 100:
                        src_seg = segmenter(remove_stopwords(clean_text(contenu)), taille)
                        res = detect_plagiarism(sus_seg, src_seg, threshold=seuil)
                        if res:
                            best = dict(max(res, key=lambda x: x["score"]))
                            best["fichier"], best["chemin"] = os.path.basename(f), f
                            resultats.append(best)
                            for r in res:
                                if r.get("suspicious"):
                                    flagged.add(r["suspicious"])
                except Exception:
                    erreurs += 1
                self.file_q.put(("progress", i + 1, total, len(resultats)))
            self.file_q.put(("fin", dict(resultats=resultats, total=total, nb_seg=len(sus_seg),
                                         flagged=len(flagged), interrompu=interrompu,
                                         erreurs=erreurs, temps=time.time() - self.t_debut,
                                         seuil=seuil)))
        except Exception as e:
            self.file_q.put(("erreur", f"Erreur pendant l'analyse : {e}"))

    def _poll(self):
        fin = None
        try:
            while True:
                msg = self.file_q.get_nowait()
                if msg[0] == "progress":
                    _, i, total, nb = msg
                    self.progress.set(i / total)
                    self.lbl_pct.configure(text=f"{int(i / total * 100)}%")
                    self.val_fichiers.configure(text=f"{i} / {total}")
                    self.val_plagiats.configure(text=str(nb))
                else:
                    fin = msg
        except queue.Empty:
            pass
        self.val_temps.configure(text=fmt_temps(time.time() - self.t_debut))
        if fin:
            self._terminer(fin)
        else:
            self.after(100, self._poll)

    def _terminer(self, msg):
        self.en_cours = False
        self.btn_lancer.configure(state="normal", text="▶  Lancer la Détection sur le Corpus",
                                  fg_color=BLUE, hover_color=BLUE_HOVER)
        self._etat_boutons(True)

        if msg[0] == "erreur":
            self.lbl_etat.configure(text="Analyse impossible")
            self.lbl_etat_sub.configure(text=msg[1])
            self.badge.configure(text="Erreur", fg_color="#fdecea", text_color=RED)
            self._set_statut("Erreur", RED)
            messagebox.showerror("Erreur", msg[1])
            return

        d = msg[1]
        d["document"] = self.doc_analyse
        resultats = sorted(d["resultats"], key=lambda x: x["score"], reverse=True)
        d["resultats"] = resultats
        couverture = min(100, d["flagged"] / max(1, d["nb_seg"]) * 100)
        self.val_compar.configure(text=f"{couverture:.0f}%")
        self.val_temps.configure(text=fmt_temps(d["temps"]))
        self.lbl_etat.configure(text="Analyse interrompue" if d["interrompu"] else "Analyse terminée")
        sub = f"{d['total']} fichier(s) comparé(s) en {d['temps']:.1f}s"
        if d["erreurs"]:
            sub += f" · {d['erreurs']} fichier(s) illisible(s)"
        self.lbl_etat_sub.configure(text=sub)
        self._set_statut("Terminé" if not d["interrompu"] else "Interrompu")

        if resultats:
            self.badge.configure(text=f"{len(resultats)} correspondance(s)", fg_color="#fdecea", text_color=RED)
        else:
            self.badge.configure(text="Aucun plagiat", fg_color="#e6f6ec", text_color=GREEN)

        self._montrer_placeholder(False)
        self.txt_resultats.delete("1.0", "end")
        self.txt_resultats.insert("1.0", self._construire_rapport(d, 10, 150))

        # fichiers similaires
        for w in self.liste_sim.winfo_children():
            w.destroy()
        if not resultats:
            ctk.CTkLabel(self.liste_sim, text="✅ Aucun fichier similaire trouvé", text_color=GREEN).pack(pady=14)
        for r in resultats[:15]:
            ligne = ctk.CTkFrame(self.liste_sim, fg_color="#f7f9fc", border_width=1, border_color=BORDER)
            ligne.pack(fill="x", pady=3, padx=2)
            nom = r["fichier"] if len(r["fichier"]) < 28 else r["fichier"][:25] + "..."
            ctk.CTkLabel(ligne, text="📄 " + nom, text_color=TEXT, font=("Arial", 12)).pack(side="left", padx=8, pady=6)
            ctk.CTkButton(ligne, text="Ouvrir", width=60, height=24,
                          command=lambda c=r["chemin"]: self.ouvrir_fichier(c)).pack(side="right", padx=6)
            ctk.CTkLabel(ligne, text=f"{r['score'] * 100:.0f}%", text_color="white", fg_color=RED,
                         corner_radius=8, width=48, height=22, font=("Arial", 11, "bold")).pack(side="right")

        # rapport détaillé + historique
        rapport = self._construire_rapport(d, 50, 300)
        self._voir_rapport_sans_changer_page(rapport)
        self.historique.append(dict(
            date=datetime.datetime.now().strftime("%d/%m/%Y %H:%M"), fichiers=d["total"],
            plagiats=len(resultats), temps=fmt_temps(d["temps"]), document=d["document"],
            extrait=self.contenu_suspect.strip()[:60].replace("\n", " "), rapport=rapport))
        self._sauver_historique()

    def _voir_rapport_sans_changer_page(self, rapport):
        self.txt_rapport.delete("1.0", "end")
        self.txt_rapport.insert("1.0", rapport)

    def _construire_rapport(self, d, limite, nb_car):
        res = d["resultats"]
        lignes = [f"Document suspect : {d.get('document', 'texte saisi')}",
                  f"Fichiers analysés : {d['total']}  |  Seuil : {d['seuil']}  |  Temps : {d['temps']:.1f}s"]
        if d["interrompu"]:
            lignes.append("(analyse interrompue par l'utilisateur)")
        if not res:
            lignes.append("\n✅ Aucun plagiat détecté dans le corpus.")
            return "\n".join(lignes)
        lignes.append(f"\n⚠️ {len(res)} correspondance(s) trouvée(s) !\n")
        for i, r in enumerate(res[:limite], 1):
            lignes.append(f"--- Match {i} | Score : {r['score']:.2f} ---")
            lignes.append(f"Fichier source : {r['fichier']}")
            lignes.append(f"Passage suspect : {str(r.get('suspicious', ''))[:nb_car]}...")
            if r.get("source"):
                lignes.append(f"Passage source  : {str(r['source'])[:nb_car]}...")
            lignes.append("")
        return "\n".join(lignes)


if __name__ == "__main__":
    app = AppModerne()
    app.mainloop()